"""文档解析（上传文件 → 纯文本 + 元数据）。

支持格式（``KB_DOCUMENT_EXTENSIONS``）：

- 纯文本类：``.md`` / ``.markdown`` / ``.txt`` / ``.csv`` / ``.json`` / ``.html`` / ``.htm``；
- 二进制类：``.pdf``（``pypdf``）、``.docx``（``python-docx``）。

未知扩展名：先按 UTF-8 尽力解码；解码失败则给出**明确错误**并列出支持的格式，
绝不静默丢弃内容。可选依赖缺失时同样给出可执行的安装指引，不吞异常。
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any

from ..constants import KB_DOCUMENT_EXTENSIONS
from ..errors import bad_request

#: HTML 标签与常见实体（无 bs4 依赖时的轻量清洗；仅用于「纯文本化」，不做结构还原）
_TAG = re.compile(r"<[^>]+>")
_SCRIPT = re.compile(r"(?is)<(script|style)[^>]*>.*?</\1>")
_BLOCK_BREAK = re.compile(r"(?i)</?(p|div|br|li|tr|h[1-6])[^>]*>")
_ENTITY = {
    "&nbsp;": " ",
    "&amp;": "&",
    "&lt;": "<",
    "&gt;": ">",
    "&quot;": '"',
    "&#39;": "'",
}
_BLANK_RUN = re.compile(r"\n{3,}")


@dataclass
class LoadedDocument:
    """解析产物：正文文本 + 展示用元数据。"""

    text: str
    meta: dict[str, Any] = field(default_factory=dict)


def _suffix(filename: str) -> str:
    return PurePosixPath(filename or "").suffix.lower()


def load_document(filename: str, data: bytes) -> LoadedDocument:
    """把上传文件解析为文本。

    :param filename: 原始文件名（用于判定扩展名与展示标题）。
    :param data: 文件字节内容。
    :raises ApiError: 格式不受支持 / 解码失败 / 解析依赖缺失 / 内容为空。
    """
    if not data:
        raise bad_request("文件内容为空，无法建立知识库文档")
    suffix = _suffix(filename)
    meta: dict[str, Any] = {"filename": filename, "extension": suffix, "sizeBytes": len(data)}

    if suffix == ".pdf":
        text, extra = _load_pdf(data)
    elif suffix == ".docx":
        text, extra = _load_docx(data)
    elif suffix in (".html", ".htm"):
        text = _strip_html(_decode_text(data, filename))
        extra = {}
    elif suffix == ".json":
        text = _prettify_json(_decode_text(data, filename))
        extra = {}
    elif suffix in (".md", ".markdown", ".txt", ".csv"):
        text = _decode_text(data, filename)
        extra = {}
    else:
        # 未知扩展名：尽力按 UTF-8 文本解析，失败则明确报错
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            supported = " / ".join(KB_DOCUMENT_EXTENSIONS)
            raise bad_request(f"无法解析文件「{filename}」：{exc}；支持的格式：{supported}") from exc
        extra = {}

    meta.update(extra)
    normalized = _BLANK_RUN.sub("\n\n", (text or "").replace("\r\n", "\n")).strip()
    if not normalized:
        raise bad_request(f"文件「{filename}」未解析出可用文本内容（可能是扫描件或空文档）")
    meta["charCount"] = len(normalized)
    return LoadedDocument(text=normalized, meta=meta)


# ------------------------------- 各格式实现 -------------------------------


def _decode_text(data: bytes, filename: str) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            # 中文环境常见回退编码
            return data.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise bad_request(f"文件「{filename}」文本编码无法识别（非 UTF-8 / GB18030）：{exc}") from exc


def _load_pdf(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - 取决于运行环境
        raise bad_request("未安装 PDF 解析依赖 pypdf，请执行 `pip install pypdf` 后重试") from exc
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - 损坏文件统一转业务错误
        raise bad_request(f"PDF 解析失败（文件可能损坏或加密）：{exc}") from exc
    pages: list[str] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            content = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - 单页异常不应导致整篇失败
            content = ""
        if content.strip():
            pages.append(f"[第 {number} 页]\n{content.strip()}")
    return "\n\n".join(pages), {"pageCount": len(reader.pages)}


def _load_docx(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - 取决于运行环境
        raise bad_request("未安装 DOCX 解析依赖 python-docx，请执行 `pip install python-docx` 后重试") from exc
    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise bad_request(f"DOCX 解析失败（文件可能损坏）：{exc}") from exc

    blocks: list[str] = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks), {"paragraphCount": len(document.paragraphs), "tableCount": len(document.tables)}


def _strip_html(html: str) -> str:
    """轻量 HTML 纯文本化：去脚本样式、块级标签转换行、去标签与实体。"""
    cleaned = _SCRIPT.sub(" ", html)
    cleaned = _BLOCK_BREAK.sub("\n", cleaned)
    cleaned = _TAG.sub("", cleaned)
    for entity, char in _ENTITY.items():
        cleaned = cleaned.replace(entity, char)
    return cleaned


def _prettify_json(raw: str) -> str:
    """JSON 展平为可读文本（对象保留键路径，数组逐项展开），便于向量检索命中字段值。"""
    try:
        parsed = json.loads(raw)
    except ValueError:
        return raw
    lines: list[str] = []
    _flatten_json(parsed, "", lines)
    return "\n".join(lines) if lines else raw


def _flatten_json(node: Any, prefix: str, lines: list[str]) -> None:
    if isinstance(node, dict):
        for key, value in node.items():
            _flatten_json(value, f"{prefix}.{key}" if prefix else str(key), lines)
    elif isinstance(node, list):
        for index, value in enumerate(node):
            _flatten_json(value, f"{prefix}[{index}]", lines)
    else:
        lines.append(f"{prefix}: {node}")


__all__ = ["LoadedDocument", "load_document"]
